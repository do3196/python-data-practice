from openpyxl import load_workbook

workbook = load_workbook("students.xlsx")
worksheet = workbook["학생점수"]


def search_student(worksheet, target_name):

    scores = []

    print(f"\n[{target_name} 학생 성적]")

    for row in range(2, worksheet.max_row + 1):

        student_id = worksheet.cell(
            row=row,
            column=1
        ).value

        name = worksheet.cell(
            row=row,
            column=2
        ).value

        subject = worksheet.cell(
            row=row,
            column=3
        ).value

        score = worksheet.cell(
            row=row,
            column=4
        ).value

        grade = worksheet.cell(
            row=row,
            column=5
        ).value


        if name == target_name:

            print(
                f"{subject} : "
                f"{score}점 / {grade}등급"
            )

            scores.append(score)


    if len(scores) > 0:

        total = sum(scores)
        average = total / len(scores)

        print(f"\n성적 개수 : {len(scores)}개")
        print(f"총점 : {total}점")
        print(f"평균 : {average:.2f}점")

    else:

        print("해당 학생이 없습니다.")


name = input("조회할 학생 이름: ")

search_student(
    worksheet,
    name
)