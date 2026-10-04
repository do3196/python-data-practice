from openpyxl import Workbook, load_workbook
from pathlib import Path

FILE_NAME = "students.xlsx"

# Excel 파일이 없으면 새로 만들기
if not Path(FILE_NAME).exists():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "학생점수"
    worksheet.append(["이름", "점수"])
    worksheet.append(["김민수", 80])
    worksheet.append(["이영희", 90])
    worksheet.append(["박철수", 85])
    workbook.save(FILE_NAME)

    print("새로운 Excel 파일을 만들었습니다.")

# 기존 Excel 파일 읽기
workbook = load_workbook(FILE_NAME)
worksheet = workbook["학생점수"]

print("\n[현재 학생 점수]")

for row in worksheet.iter_rows(min_row=2, values_only=True):
    name = row[0]
    score = row[1]

    print(f"{name} : {score}점")

# 새로운 학생 입력
new_name = input("\n추가할 학생 이름: ")

try:
    new_score = int(input("점수: "))
    # 새로운 행 추가
    worksheet.append([new_name, new_score])
    # Excel에 저장
    workbook.save(FILE_NAME)
    print("새로운 학생 정보가 저장되었습니다.")

except ValueError:
    print("점수는 숫자로 입력하세요.")