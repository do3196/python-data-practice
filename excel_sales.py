# openpyxl 라이브러리에서 Excel 파일을 불러오는 load_workbook 기능을 가져옴.
from openpyxl import load_workbook

# 이 Python 파일과 같은 폴더에 sales.xlsx 파일이 있어야 함.
INPUT_FILE = "sales.xlsx"

# 처리 결과를 저장할 새 Excel 파일 이름.
OUTPUT_FILE = "sales_result.xlsx"

# Excel 파일 안에서 사용할 시트(탭)의 이름.
SHEET_NAME = "매출"

# Excel 파일 불러오기
# INPUT_FILE에 지정한 Excel 파일을 열어 workbook 변수에 저장
# workbook은 Excel 파일 전체를 의미
workbook = load_workbook(INPUT_FILE)

# 지정한 이름의 시트가 Excel 파일 안에 있는지 확인.
# workbook.sheetnames는 현재 Excel 파일에 있는 모든 시트 이름 목록.
if SHEET_NAME not in workbook.sheetnames:
    # "매출" 시트가 없다면 오류를 발생시키고 프로그램을 중단.
    raise ValueError(f"'{SHEET_NAME}' 시트가 없습니다.")

# "매출"이라는 이름의 시트를 선택하여 worksheet 변수에 저장.
# worksheet는 Excel 파일의 특정 시트 한 장을 의미.
worksheet = workbook[SHEET_NAME]

# D열에 실제/예상 구분 추가
# D열 1행, 즉 D1 셀에 "구분"이라는 제목을 입력.
# D열에는 "실제" 또는 "예상"이라는 값을 저장.
worksheet["D1"] = "구분"

# Excel 데이터를 sales 리스트로 가져오기
# Excel에서 읽은 매출 데이터를 저장할 빈 리스트를 만듬.
# 각 항목은 [연도, 월, 매출액, 구분] 형태로 저장.
sales = []

# Excel의 2행부터 마지막 행까지 반복. 1행은 제목 행이므로 2행부터 시작.
# worksheet.max_row는 현재 시트에서 데이터가 있는 마지막 행 번호.
for row in range(2, worksheet.max_row + 1):
    # 현재 행의 A열(1번째 열)에서 연도 값을 읽음.
    year = worksheet.cell(row=row, column=1).value
    # 현재 행의 B열(2번째 열)에서 월 값을 읽음.
    month = worksheet.cell(row=row, column=2).value
    # 현재 행의 C열(3번째 열)에서 매출 값을 읽음.
    value = worksheet.cell(row=row, column=3).value
    # 현재 행의 D열(4번째 열)에서 "실제" 또는 "예상" 구분 값을 읽음.
    data_type = worksheet.cell(row=row, column=4).value
    # 연도, 월, 매출 중 하나라도 비어 있으면 정상 데이터가 아니므로 건너뜀.
    if year is None or month is None or value is None:
        continue

    # D열 구분값이 비어 있다면, 기존 데이터는 실제 매출이라고 판단함.
    if data_type is None:
        # Python 변수에 "실제"라는 값을 저장합니다.
        data_type = "실제"
        # Excel 시트의 D열에도 "실제"라는 값을 입력합니다.
        worksheet.cell(row=row, column=4).value = data_type

    # 현재 행의 데이터를 sales 리스트에 추가합니다.
    # int(year): 연도를 정수형으로 변환합니다. 예: 2025.0 → 2025
    # int(month): 월을 정수형으로 변환합니다. 예: 4.0 → 4
    # float(value): 매출을 실수형으로 변환합니다. 예: 128 → 128.0
    # data_type: "실제" 또는 "예상" 값을 저장합니다.
    sales.append([
        int(year),
        int(month),
        float(value),
        data_type
    ])

# 기존 매출 수정 함수
# 특정 연도와 월의 매출을 수정하는 함수를 정의합니다.
# year: 수정할 연도, month: 수정할 월, new_value: 새로 입력할 매출액
def update_sale(year, month, new_value):
    """
    Excel과 sales 리스트의 특정 연월 매출을 수정합니다.
    """

    # 해당 연도와 월의 데이터가 실제로 존재하는지 확인하기 위한 변수임.
    # 처음에는 찾지 못했으므로 False로 설정.
    found = False

    # Excel 시트의 2행부터 마지막 행까지 반복하며 원하는 연도와 월을 찾음.
    for row in range(2, worksheet.max_row + 1):

        # 현재 행 A열에서 연도 값을 읽음.
        excel_year = worksheet.cell(row=row, column=1).value
        # 현재 행 B열에서 월 값을 읽음.
        excel_month = worksheet.cell(row=row, column=2).value
        # 현재 행의 연도와 월이 수정하려는 연도와 월과 같은지 확인.
        if excel_year == year and excel_month == month:
            # 찾은 행의 C열 매출 값을 new_value로 변경.
            worksheet.cell(row=row, column=3).value = new_value
            # 수정한 값은 실제 매출이므로 D열에 "실제"라고 기록.
            worksheet.cell(row=row, column=4).value = "실제"
            # 원하는 데이터를 찾았다는 표시로 found를 True로 변경.
            found = True
            # 원하는 행을 수정했으므로 Excel 반복문을 종료.
            break
    # Python의 sales 리스트에서도 같은 연도와 월의 데이터를 찾음.
    for item in sales:
        # item[0]은 연도, item[1]은 월입니다.
        # 둘 다 수정하려는 year, month와 같으면 해당 데이터를 수정.
        if item[0] == year and item[1] == month:
            # item[2]는 매출액이므로 새 매출액을 실수형으로 변환하여 저장.
            item[2] = float(new_value)
            # item[3]은 구분값이므로 "실제"로 변경.
            item[3] = "실제"
            # 원하는 데이터를 수정했으므로 리스트 반복문을 종료.
            break
    # 데이터를 찾고 수정했으면 True, 찾지 못했으면 False를 반환.
    return found

# 특정 연도의 월별 매출 가져오기
# 지정한 연도의 1월부터 12월까지 매출액을 가져오는 함수를 정의.
# 예: get_year_sales(2025)
def get_year_sales(year):
    """
    sales에서 원하는 연도의 1월~12월 데이터를 가져옵니다.
    """

    # 지정 연도의 데이터를 임시로 저장할 빈 리스트.
    year_data = []

    # sales 리스트에 있는 모든 매출 데이터를 하나씩 확인.
    for item in sales:
        # item[0]은 연도.
        # 현재 데이터의 연도가 찾고 싶은 year와 같으면 리스트에 추가.
        if item[0] == year:
            year_data.append(item)

    # year_data를 월(item[1]) 기준으로 오름차순 정렬.
    # 즉, 1월, 2월, 3월 ... 12월 순서로 정렬.
    year_data.sort(key=lambda item: item[1])

    # 해당 연도에 실제로 어떤 월 데이터가 있는지 저장할 빈 리스트.
    months = []

    # 정렬된 연도별 데이터를 하나씩 확인.
    for item in year_data:
        # item[1]은 월이므로 months 리스트에 월 번호를 추가.
        months.append(item[1])

    # months가 정확히 [1, 2, 3, ..., 12]인지 확인.
    # list(range(1, 13))은 1부터 12까지의 리스트를 만듬.
    if months != list(range(1, 13)):
        # 1~12월 중 하나라도 빠졌거나 중복 또는 잘못된 월이 있다면 오류를 발생.
        raise ValueError(
            f"{year}년의 1월~12월 데이터가 모두 필요합니다."
        )

    # 1월~12월의 매출액만 따로 저장할 빈 리스트.
    values = []

    # year_data의 데이터를 하나씩 확인.
    for item in year_data:

        # item[2]는 매출액이므로 values 리스트에 추가.
        values.append(item[2])

    # 1월부터 12월까지의 매출액 리스트를 반환.
    # 예: [100.0, 110.0, 120.0, ..., 200.0]
    return values

# 평균 변화량 계산
# 월별 매출 데이터에서 전월 대비 변화량의 평균을 계산하는 함수를 정의.
# data에는 보통 12개월 매출액이 들어옴.
def calculate_average_change(data):
    """
    월별 변화량을 구한 후 평균을 반환합니다.
    """

    # 전월 대비 변화량을 저장할 빈 리스트입니다.
    changes = []

    # 두 번째 값부터 마지막 값까지 반복합니다.
    # i=1은 두 번째 달이며, 첫 번째 달과 비교할 수 있습니다.
    for i in range(1, len(data)):

        # 현재 달 매출에서 이전 달 매출을 빼서 변화량을 계산합니다.
        # 예: 2월 120 - 1월 100 = +20
        change = data[i] - data[i - 1]

        # 계산한 변화량을 changes 리스트에 추가합니다.
        changes.append(change)

    # 모든 월별 변화량의 합계를 변화량 개수로 나누어 평균 변화량을 구합니다.
    # 12개월 데이터라면 변화량은 11개입니다.
    average_change = sum(changes) / len(changes)

    # 계산한 평균 변화량을 함수 밖으로 반환합니다.
    return average_change

# 다음 연도 매출을 예측하고 Excel에 추가
# 기준 연도(base_year)의 데이터를 사용하여 다음 연도 매출을 예측하는 함수를 정의합니다.
# 예: base_year가 2025이면 2026년 예상 매출을 만듭니다.
def predict_and_append_year(base_year):
    """
    base_year 데이터를 이용해 다음 연도의
    1월~12월 예상 매출을 계산하고 저장합니다.
    """

    # 예측 대상 연도는 기준 연도보다 1년 뒤입니다.
    # 예: 2025 + 1 = 2026
    target_year = base_year + 1

    # sales 리스트에 이미 목표 연도 데이터가 있는지 확인합니다.
    for item in sales:
        # 목표 연도 데이터가 하나라도 발견되면 중복 생성을 막습니다.
        if item[0] == target_year:
            # 이미 데이터가 있다는 오류 메시지를 표시하고 프로그램을 중단합니다.
            raise ValueError(
                f"{target_year}년 데이터가 이미 존재합니다."
            )

    # 기준 연도의 1월~12월 매출을 가져옵니다.
    # 예: 2025년 12개월 매출액 리스트
    base_sales = get_year_sales(base_year)

    # 기준 연도의 월별 변화량 평균을 계산합니다.
    average_change = calculate_average_change(base_sales)

    # 기준 연도 12월의 매출액을 가져옵니다.
    # base_sales[-1]은 리스트의 마지막 값, 즉 12월 매출입니다.
    current_value = base_sales[-1]

    # 계산한 평균 변화량을 화면에 소수점 둘째 자리까지 출력합니다.
    print(
        f"\n{base_year}년 평균 변화량: "
        f"{average_change:.2f}"
    )
    # 1월부터 12월까지 반복하여 다음 연도의 예상 매출을 생성합니다.
    for month in range(1, 13):
        # 이전 월 매출에 평균 변화량을 더해 이번 달의 예상값을 계산합니다.
        current_value = current_value + average_change
        # 계산된 예상 매출을 반올림하여 정수로 만듭니다.
        # 예: 131.7 → 132
        predicted_value = round(current_value)
        # 새 데이터 한 행을 만듭니다.
        # 순서: [연도, 월, 예상 매출, 구분]
        new_data = [target_year, month, predicted_value, "예상"]

        # 새 예상 데이터를 Python의 sales 리스트에 추가합니다.
        sales.append(new_data)

        # 새 예상 데이터를 Excel 시트의 마지막 행 아래에 추가합니다.
        # A열부터 D열까지 순서대로 기록됩니다.
        worksheet.append(new_data)

        # 다음 달 예측은 반올림된 이번 달 예상값을 기준으로 계산합니다.
        # 따라서 Excel에 저장된 값과 다음 달 계산 기준이 일치합니다.
        current_value = predicted_value

        # 현재 생성한 월별 예상 매출을 화면에 출력합니다.
        print(
            f"{target_year}년 {month}월 "
            f"예상 매출: {predicted_value}"
        )


# ---------------------------------------------
# 기존 실제 매출 수정 예제
# ---------------------------------------------
# 아래 줄은 2025년 4월 매출을 130으로 수정하는 예시입니다.
# 현재는 줄 맨 앞에 #이 있으므로 실행되지 않습니다.
# 필요하다면 맨 앞의 #을 삭제하여 실행합니다.
# update_sale(2025, 4, 130)
# ---------------------------------------------
# 2026년 예상 매출 추가
# ---------------------------------------------
# 2025년 실제 매출 데이터를 바탕으로 2026년 1월~12월 예상 매출을 생성합니다.
# 생성된 데이터는 sales 리스트와 Excel 시트에 함께 추가됩니다.
predict_and_append_year(2025)
# ---------------------------------------------
# 2026년 예상 매출을 이용해 2027년 예상 매출 추가
# ---------------------------------------------

# 바로 앞에서 생성한 2026년 예상 매출을 바탕으로
# 2027년 1월~12월 예상 매출을 추가로 생성합니다.
predict_and_append_year(2026)
# ---------------------------------------------
# 새로운 Excel 파일로 저장
# ---------------------------------------------

# 수정 및 추가된 내용을 OUTPUT_FILE 이름으로 저장합니다.
# 원본 sales.xlsx 파일은 그대로 두고, 새 파일이 만들어집니다.
workbook.save(OUTPUT_FILE)

# Excel 파일 저장이 완료되었음을 화면에 출력합니다.
print("\nExcel 저장이 완료되었습니다.")

# 실제 저장한 파일 이름을 화면에 출력합니다.
print(f"저장 파일: {OUTPUT_FILE}")

