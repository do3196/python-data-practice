from openpyxl import load_workbook

workbook = load_workbook("studentsGradeData.xlsx")
worksheet = workbook["학생점수"]

def analyze_student(worksheet, target_name):
    student_data = []
    # 학생 데이터 검색
    for row in range(2, worksheet.max_row + 1):
        name = worksheet.cell(row=row, column=2).value
        subject = worksheet.cell(row=row, column=3).value
        score = worksheet.cell(row=row, column=4).value
        grade = worksheet.cell(row=row, column=5).value

        if name == target_name:
            student_data.append([subject, score, grade])
    # 학생이 없는 경우
    if len(student_data) == 0:
        print("해당 학생이 없습니다.")
        return

    print(f"\n[{target_name} 학생 성적]")

    scores = []
    # 과목별 성적 출력
    for data in student_data:
        subject = data[0]
        score = data[1]
        grade = data[2]
        print(f"{subject} : " f"{score}점 / {grade}등급")
        scores.append(score)

    # 총점과 평균
    total = sum(scores)
    average = total / len(scores)

    print(f"\n총점 : {total}점")
    print(f"평균 : {average:.2f}점")

    # 평균과 비교
    print("\n[평균 비교]")

    for data in student_data:
        subject = data[0]
        score = data[1]
        if score >= average:
            result = "평균 이상"
        else:
            result = "평균 미만"
        print(f"{subject} : {result}")

    # 최고 / 최저 점수
    max_score = max(scores)
    min_score = min(scores)

    # 최고 / 최저 과목 찾기
    max_subject = ""
    min_subject = ""

    for data in student_data:
        subject = data[0]
        score = data[1]
        if score == max_score:
            max_subject = subject

        if score == min_score:
            min_subject = subject

    print(f"\n최고 점수 과목 : "
        f"{max_subject} / {max_score}점")

    print(f"최저 점수 과목 : "
        f"{min_subject} / {min_score}점")


name = input("조회할 학생 이름: ")

analyze_student(worksheet, name)