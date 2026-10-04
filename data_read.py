from openpyxl import load_workbook

workbook = load_workbook("students.xlsx")
worksheet = workbook["학생점수"]

def show_students(worksheet):
    print("\n[학생 점수 목록]")
    student_count = 0

    for row in range(2, worksheet.max_row + 1):
        name = worksheet.cell(row=row, column=1).value
        score = worksheet.cell(row=row, column=2).value

        # 빈 행은 출력하지 않음
        if name is None:
            continue

        print(f"{name} : {score}점")
        student_count += 1

    if student_count == 0:
        print("등록된 학생이 없습니다.")

def find_student_row(worksheet, target_name=None):
    # 메뉴 2에서 호출한 경우 이름을 직접 입력받음
    search_menu = target_name is None

    if search_menu:
        target_name = input("검색할 학생 이름: ").strip()

        if target_name == "":
            print("학생 이름을 입력하세요.")
            return None

    # Excel의 학생 이름을 하나씩 확인
    for row in range(2, worksheet.max_row + 1):
        name = worksheet.cell(row=row, column=1).value

        if name == target_name:

            # 메뉴 2에서 검색한 경우 점수까지 출력
            if search_menu:
                score = worksheet.cell(row=row, column=2).value
                print(f"{target_name} 학생의 점수는 "
                      f"{score}점입니다."
                )
            # 학생이 저장된 행 번호 반환
            return row
    # 검색 메뉴에서 학생을 찾지 못한 경우
    if search_menu:
        print("해당 학생이 없습니다.")

    return None

def add_student(workbook, worksheet,  name=None, score=None):
 # 메뉴 3에서 호출하면 이름을 직접 입력받음
    if name is None:
        name = input("추가할 학생 이름: ").strip()

    if name == "":
        print("학생 이름을 입력하세요.")
        return
    # 중복 학생 확인
    row_number = find_student_row(worksheet, name)

    if row_number is not None:
        print("이미 등록된 학생입니다.")
        return

    # 점수가 전달되지 않았다면 사용자에게 입력받음
    if score is None:
        score = input("점수를 입력하세요: ")

    try:
        score = int(score)
        # 점수 범위 확인
        if score < 0 or score > 100:
            print("점수는 0~100 사이로 입력하세요.")
            return
    except ValueError:
        print("점수는 숫자로 입력하세요.")
        return

    # Excel 마지막 행에 학생 정보 추가
    worksheet.append([name, score])
    # 수정한 Excel 파일 저장
    workbook.save("students.xlsx")
    print("학생 정보가 저장되었습니다.")


# 메뉴 프로그램
while True:
    print("\n==========================")
    print("      학생 점수 관리")
    print("==========================")
    print("1. 전체 목록 보기")
    print("2. 학생 검색하기")
    print("3. 학생 추가하기")
    print("0. 종료하기")

    menu = input("메뉴를 선택하세요: ").strip()

    if menu == "1":
        show_students(worksheet)
    elif menu == "2":
        find_student_row(worksheet)
    elif menu == "3":
        add_student(workbook, worksheet)
    elif menu == "0":
        print("프로그램을 종료합니다.")
        workbook.close()
        break

    else:
        print("0, 1, 2, 3 중에서 선택하세요.")