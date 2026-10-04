from openpyxl import load_workbook

# 기존 Excel 파일 열기
workbook = load_workbook("students.xlsx")

# 학생점수 시트 선택
worksheet = workbook["학생점수"]

print("[학생 점수 목록]")

# 두 번째 행부터 마지막 행까지 읽기
for line in worksheet.iter_rows(min_row=2, values_only=True):
    name = line[0]
    score = line[1]
    print(f"이름: {name}, 점수: {score}")